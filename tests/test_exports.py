import csv
import io

from openpyxl import load_workbook

from app.utils.exports import export_csv, export_excel


def test_csv_export_neutralizes_formula_prefixed_cells():
    resp = export_csv("report", ["Name", "Notes"], [
        ["=HYPERLINK(\"http://evil.example/\",\"click\")", "safe text"],
        ["+1+1", "-2+3"],
        ["@SUM(A1)", "\tstill dangerous"],
    ])
    reader = csv.reader(io.StringIO(resp.get_data(as_text=True)))
    rows = list(reader)
    assert rows[1][0] == "'=HYPERLINK(\"http://evil.example/\",\"click\")"
    assert rows[2][0] == "'+1+1"
    assert rows[2][1] == "'-2+3"
    assert rows[3][0] == "'@SUM(A1)"
    assert rows[3][1] == "'\tstill dangerous"
    assert rows[1][1] == "safe text"  # untouched - doesn't start with a formula prefix


def test_excel_export_neutralizes_formula_prefixed_cells():
    resp = export_excel("report", ["Name"], [["=cmd|'/c calc'!A0"]])
    wb = load_workbook(io.BytesIO(resp.get_data()))
    ws = wb.active
    assert ws["A2"].value == "'=cmd|'/c calc'!A0"


def test_csv_export_leaves_numeric_and_plain_cells_untouched():
    resp = export_csv("report", ["Count", "Name"], [[42, "Regular Vendor"]])
    text = resp.get_data(as_text=True)
    assert "42" in text
    assert "Regular Vendor" in text
    assert "'42" not in text
